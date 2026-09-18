#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
成果輸出 —— 由 data/ 產生可交付的報表與圖資
    python tools/export.py [年度id]      預設取 manifest 的 defaultYear

輸出到 out/：
    <年度>年度巡檢覆蓋率統計報表.xlsx     行政區統計、巡檢次數、待補缺口、計算說明
    <年度>年度各行政區巡檢長度表.xlsx     各區 × 1-12 月（不重複／含重複／重複度對照）
    <年度>年度巡檢覆蓋成果_管段明細.geojson  可直接匯入 QGIS／ArcGIS
"""
import json, sys
from collections import defaultdict
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import ColorScaleRule

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / 'out'; OUT.mkdir(exist_ok=True)
MF = json.loads((ROOT / 'data' / 'manifest.json').read_text(encoding='utf-8'))
YID = sys.argv[1] if len(sys.argv) > 1 else MF['defaultYear']
YM = next(y for y in MF['years'] if y['id'] == YID)
NET = json.loads((ROOT / 'data' / MF['networks'][YM['network']]['file']).read_text(encoding='utf-8'))
YR = json.loads((ROOT / 'data' / YM['file']).read_text(encoding='utf-8'))

U, D = NET['units'], NET['dists']
N = len(U['len'])
STN = ['已巡檢', '列管人孔未開孔', '無人孔(虛接點/分岔)', '端點編號缺漏', '實人孔未列管']

def status(i):
    b = U['blocked'][i]
    if b == 1: return 2
    if b == 2: return 3
    if YR['cnt'][i] > 0: return 0
    return 1 if YR['listed'][i] else 4

ST = [status(i) for i in range(N)]
TOT = sum(U['len'])

F = 'Microsoft JhengHei'
HDR = PatternFill('solid', fgColor='1A2D3A'); SUB = PatternFill('solid', fgColor='2A4355')
ALT = PatternFill('solid', fgColor='F2F6F9'); TOTF = PatternFill('solid', fgColor='DCE8F0')
TH = Side(style='thin', color='C8D4DC'); MED = Side(style='medium', color='8CA3B3')
BD = Border(left=TH, right=TH, top=TH, bottom=TH)


def head(ws, cols, widths, row=1, fill2=None):
    for j, (c, w) in enumerate(zip(cols, widths), 1):
        cell = ws.cell(row, j, c)
        cell.font = Font(F, 10, bold=True, color='FFFFFF')
        cell.fill = fill2 if (fill2 and 1 < j <= 13) else HDR
        cell.border = BD
        cell.alignment = Alignment('center', 'center', wrap_text=True)
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[row].height = 30
    ws.freeze_panes = ws.cell(row + 1, 2)


# ══════════════ 彙總 ══════════════
dist_tot = defaultdict(float); dist_st = defaultdict(lambda: [0.0] * 5)
dist_cnt = defaultdict(lambda: [0.0] * 4)
dup = defaultdict(lambda: [0.0] * 13); uni = defaultdict(lambda: [0.0] * 13)
for i in range(N):
    d, L = U['d'][i], U['len'][i]
    dist_tot[d] += L; dist_st[d][ST[i]] += L
    dist_cnt[d][min(YR['cnt'][i], 3)] += L
    m = YR['m'][i]
    first = 0
    for k in range(12):
        if m >> k & 1:
            dup[d][k + 1] += L
            if not first: first = k + 1
    if first: uni[d][first] += L
order = sorted(D, key=lambda x: -(dist_tot[D.index(x)] - dist_st[D.index(x)][0]))

# ══════════════ 報表一：覆蓋率統計 ══════════════
wb = Workbook(); wb.remove(wb.active)
ws = wb.create_sheet('行政區巡檢統計')
cols = ['行政區', '總管線長度(km)', '已巡檢長度(km)', '未巡檢長度(km)', '巡檢覆蓋率',
        '未巡檢：列管人孔未開孔', '未巡檢：兩端無人孔', '未巡檢：圖資編號缺漏',
        '未巡檢：人孔未列管', '巡檢1次(km)', '巡檢2次(km)', '巡檢3次以上(km)']
head(ws, cols, [10, 13, 13, 13, 10, 15, 14, 14, 14, 11, 11, 12])
for i, name in enumerate(order):
    d = D.index(name); rw = i + 2
    s, c = dist_st[d], dist_cnt[d]
    vals = [name, dist_tot[d] / 1000, s[0] / 1000, None, None,
            s[1] / 1000, s[2] / 1000, s[3] / 1000, s[4] / 1000,
            c[1] / 1000, c[2] / 1000, c[3] / 1000]
    for j, v in enumerate(vals, 1):
        cc = ws.cell(rw, j, v); cc.font = Font(F, 10); cc.border = BD
        if j > 1: cc.number_format = '#,##0.00'
        if i % 2: cc.fill = ALT
    ws.cell(rw, 4, f'=B{rw}-C{rw}').number_format = '#,##0.00'
    ws.cell(rw, 5, f'=IF(B{rw}=0,"",C{rw}/B{rw})').number_format = '0.0%'
    for j in (4, 5):
        ws.cell(rw, j).font = Font(F, 10, bold=(j == 5)); ws.cell(rw, j).border = BD
        if i % 2: ws.cell(rw, j).fill = ALT
n = len(order) + 1
ws.cell(n + 1, 1, '全市合計').font = Font(F, 10, bold=True)
for j in list(range(2, 5)) + list(range(6, 13)):
    L = get_column_letter(j)
    cc = ws.cell(n + 1, j, f'=SUM({L}2:{L}{n})')
    cc.font = Font(F, 10, bold=True); cc.number_format = '#,##0.00'
ws.cell(n + 1, 5, f'=C{n+1}/B{n+1}').font = Font(F, 10, bold=True)
ws.cell(n + 1, 5).number_format = '0.0%'
for j in range(1, 13):
    ws.cell(n + 1, j).fill = TOTF; ws.cell(n + 1, j).border = BD

ws2 = wb.create_sheet('巡檢次數分布')
head(ws2, ['巡檢次數', '管線長度(km)', '占全市比例', '說明'], [14, 14, 12, 46])
tot_cnt = [sum(dist_cnt[d][k] for d in dist_cnt) for k in range(4)]
notes = ['未巡檢：兩端人孔皆未開孔，或該段本就沒有人孔可開',
         '巡檢 1 次，符合年度抽樣原則', '同段重複巡檢 2 次',
         '重複度偏高，資源可考慮移轉至未巡檢區段']
for i, (lab, v, note) in enumerate(zip(['0 次', '1 次', '2 次', '3 次以上'], tot_cnt, notes)):
    rw = i + 2
    for j, val in enumerate([lab, v / 1000, None, note], 1):
        cc = ws2.cell(rw, j, val); cc.font = Font(F, 10); cc.border = BD
        if j == 2: cc.number_format = '#,##0.00'
        if j == 4: cc.alignment = Alignment(vertical='center', wrap_text=True)
    ws2.cell(rw, 3, f'=B{rw}/SUM($B$2:$B$5)').number_format = '0.0%'
    ws2.cell(rw, 3).font = Font(F, 10); ws2.cell(rw, 3).border = BD
ws2.cell(6, 1, '合計').font = Font(F, 10, bold=True)
ws2.cell(6, 2, '=SUM(B2:B5)').font = Font(F, 10, bold=True)
ws2.cell(6, 2).number_format = '#,##0.00'

ws3 = wb.create_sheet('待補缺口清單')
head(ws3, ['排序', '行政區', '管段編號', '長度(m)', '未巡檢原因', '歸屬人孔', '建議處置'],
     [6, 10, 30, 11, 22, 14, 30])
sug = {1: '排入下年度開孔計畫即可補足', 2: '無孔可開，需增設人孔或改用 CCTV 等方式',
       3: '先補正管線端點編號與人孔圖資對應', 4: '評估納入列管人孔名冊'}
gap = sorted((i for i in range(N) if ST[i] != 0), key=lambda i: -U['len'][i])[:400]
for k, i in enumerate(gap):
    rw = k + 2
    vals = [k + 1, D[U['d'][i]], U['pinum'][i][:70], round(U['len'][i], 1),
            STN[ST[i]], U['node'][i] or '—', sug.get(ST[i], '')]
    for j, v in enumerate(vals, 1):
        cc = ws3.cell(rw, j, v); cc.font = Font(F, 9.5); cc.border = BD
        if j == 4: cc.number_format = '#,##0.0'
        if k % 2: cc.fill = ALT

ws4 = wb.create_sheet('計算邏輯說明')
ws4.column_dimensions['A'].width = 22; ws4.column_dimensions['B'].width = 96
st = NET['stats']
txt = [
    ('計算原則', ''),
    ('歸屬方式', '管線上的每一點，歸屬於「沿管線走最近的那座實人孔」。該人孔若有開孔紀錄，'
                 '這段就算已巡檢。'),
    ('退化為各一半', '兩座人孔直接相連時，分界點恰好落在中點，等同於「各負責一半」的傳統算法；'
                     f'本次共有 {st.get("split", 0):,} 條管線被分界點切開。'),
    ('末端處理', '某一側沒有人孔時（管線末端），整段歸唯一的那座人孔。'),
    ('分岔處理', '分岔點依網路距離自然分界，不會把整段誤判給遠處的人孔，'
                 '也不會因為分岔就放棄整段。'),
    ('無法巡檢', '整個連通管網中沒有任何實人孔者，無論如何都無法以開孔方式巡檢。'),
    ('長度守恆', f'所有受檢單元長度總和 = 管線圖資總長度 {TOT/1000:,.2f} km，'
                 '誤差僅來自小數捨入，確保不重複計算、不遺漏。'),
    ('', ''),
    ('圖資處理', ''),
    ('拓樸建構', f'以管線屬性 US_MH／DS_MH 建立網路，{st.get("pipes", 0):,} 條管線。'),
    ('虛人孔', '人孔圖層中的「虛人」是管線轉折點，不是真實人孔，因此不列為受檢端點；'
               '分配時自然被穿越，不會被誤判成受檢人孔。'),
    ('編號吸附', f'{st.get("snapped", 0)} 個管線端點編號不存在於人孔圖層，'
                 '改以座標 1 公尺容差吸附回真實人孔。'),
    ('座標系統', 'TWD97 二度分帶（EPSG:3826），已驗證與圖資內建經緯度欄位誤差小於 1 公分。'),
]
for i, (a, b) in enumerate(txt):
    rw = i + 1
    ca, cb = ws4.cell(rw, 1, a), ws4.cell(rw, 2, b)
    bold = (b == '' and a != '')
    ca.font = Font(F, 11 if bold else 10, bold=True, color='1A2D3A' if bold else '000000')
    cb.font = Font(F, 10); cb.alignment = Alignment(vertical='top', wrap_text=True)
    ca.alignment = Alignment(vertical='top')
    if bold:
        for j in (1, 2): ws4.cell(rw, j).fill = TOTF
    ws4.row_dimensions[rw].height = 15 if bold else max(15, 15 * (len(b) // 46 + 1))
wb.save(OUT / f'{YID}年度巡檢覆蓋率統計報表.xlsx')

# ══════════════ 報表二：各區 × 月份長度表 ══════════════
wb2 = Workbook(); wb2.remove(wb2.active)

def month_sheet(name, src, note):
    ws = wb2.create_sheet(name)
    ws.merge_cells('A1:P1')
    ws['A1'] = f'{YID} 年度雨水下水道巡檢長度表（{name}）　單位：公里'
    ws['A1'].font = Font(F, 13, bold=True); ws['A1'].alignment = Alignment('center', 'center')
    ws.row_dimensions[1].height = 26
    ws.merge_cells('A2:P2'); ws['A2'] = note
    ws['A2'].font = Font(F, 9, color='5A6B78'); ws['A2'].alignment = Alignment('left', 'center')
    ws.row_dimensions[2].height = 18
    head(ws, ['行政區'] + [f'{m}月' for m in range(1, 13)] + ['總計', '總長度', '巡檢率'],
         [10] + [8.5] * 12 + [10, 10, 9], row=3, fill2=SUB)
    names = sorted(D, key=lambda x: -dist_tot[D.index(x)])
    for i, nm in enumerate(names):
        d = D.index(nm); rw = i + 4
        c = ws.cell(rw, 1, nm); c.font = Font(F, 10); c.border = BD
        if i % 2: c.fill = ALT
        for m in range(1, 13):
            v = src[d][m] / 1000
            cc = ws.cell(rw, m + 1, round(v, 3) if v > 0.0005 else None)
            cc.font = Font(F, 10); cc.border = BD; cc.number_format = '#,##0.00'
            if i % 2: cc.fill = ALT
        for j, f in [(14, f'=SUM(B{rw}:M{rw})'), (15, None), (16, f'=IF(O{rw}=0,"",N{rw}/O{rw})')]:
            cc = ws.cell(rw, j, f if f else round(dist_tot[d] / 1000, 3))
            cc.font = Font(F, 10, bold=(j != 15)); cc.border = BD
            cc.number_format = '0.0%' if j == 16 else '#,##0.00'
            if i % 2: cc.fill = ALT
    n = len(names) + 3
    ws.cell(n + 1, 1, '全市合計').font = Font(F, 10, bold=True)
    for j in range(2, 16):
        L = get_column_letter(j)
        cc = ws.cell(n + 1, j, f'=SUM({L}4:{L}{n})')
        cc.font = Font(F, 10, bold=True); cc.number_format = '#,##0.00'
    ws.cell(n + 1, 16, f'=N{n+1}/O{n+1}').font = Font(F, 10, bold=True)
    ws.cell(n + 1, 16).number_format = '0.0%'
    for j in range(1, 17):
        ws.cell(n + 1, j).fill = TOTF
        ws.cell(n + 1, j).border = Border(left=TH, right=TH, top=MED, bottom=MED)
    ws.conditional_formatting.add(f'B4:M{n}',
        ColorScaleRule(start_type='num', start_value=0, start_color='FFFFFF',
                       end_type='max', end_color='7FC7B0'))

month_sheet('不重複', uni,
  '不重複＝同一管段只計算一次，以「首次巡檢月份」歸屬。總計即為實際巡檢覆蓋長度，巡檢率＝實際覆蓋率。')
month_sheet('含重複', dup,
  '含重複＝同一管段若在多個月份被巡檢，各月皆計入，反映實際投入的工作量。'
  '此表巡檢率為「巡檢延長率」，可能超過 100%，不等於覆蓋率。')

ws = wb2.create_sheet('重複度對照')
ws.merge_cells('A1:F1'); ws['A1'] = f'{YID} 年度巡檢重複度對照　單位：公里'
ws['A1'].font = Font(F, 13, bold=True); ws['A1'].alignment = Alignment('center', 'center')
ws.row_dimensions[1].height = 26
head(ws, ['行政區', '總長度', '含重複總計', '不重複總計', '重複巡檢長度', '重複度'],
     [12, 12, 13, 13, 14, 10], row=2)
rows = sorted(D, key=lambda x: -(sum(dup[D.index(x)][1:]) - sum(uni[D.index(x)][1:])))
for i, nm in enumerate(rows):
    d = D.index(nm); rw = i + 3
    for j, v in enumerate([nm, dist_tot[d] / 1000, sum(dup[d][1:]) / 1000,
                           sum(uni[d][1:]) / 1000, None, None], 1):
        c = ws.cell(rw, j, round(v, 3) if isinstance(v, float) else v)
        c.font = Font(F, 10); c.border = BD
        if j > 1: c.number_format = '#,##0.00'
        if i % 2: c.fill = ALT
    ws.cell(rw, 5, f'=C{rw}-D{rw}').number_format = '#,##0.00'
    ws.cell(rw, 6, f'=IF(D{rw}=0,"",E{rw}/D{rw})').number_format = '0.0%'
    for j in (5, 6):
        ws.cell(rw, j).font = Font(F, 10, bold=(j == 5)); ws.cell(rw, j).border = BD
        if i % 2: ws.cell(rw, j).fill = ALT
n = len(rows) + 2
ws.cell(n + 1, 1, '全市合計').font = Font(F, 10, bold=True)
for j in range(2, 6):
    L = get_column_letter(j)
    c = ws.cell(n + 1, j, f'=SUM({L}3:{L}{n})')
    c.font = Font(F, 10, bold=True); c.number_format = '#,##0.00'
ws.cell(n + 1, 6, f'=E{n+1}/D{n+1}').font = Font(F, 10, bold=True)
ws.cell(n + 1, 6).number_format = '0.0%'
for j in range(1, 7):
    ws.cell(n + 1, j).fill = TOTF
    ws.cell(n + 1, j).border = Border(left=TH, right=TH, top=MED, bottom=MED)
wb2.save(OUT / f'{YID}年度各行政區巡檢長度表.xlsx')

# ══════════════ GeoJSON ══════════════
feats = []
for i in range(N):
    ms = [k + 1 for k in range(12) if YR['m'][i] >> k & 1]
    feats.append({'type': 'Feature', 'properties': {
        '行政區': D[U['d'][i]], '管段編號': U['pinum'][i], '歸屬人孔': U['node'][i],
        '巡檢狀態': STN[ST[i]], '是否已巡檢': '是' if ST[i] == 0 else '否',
        '巡檢次數': YR['cnt'][i], '巡檢月份': '、'.join(f'{m}月' for m in ms),
        '首次巡檢月': ms[0] if ms else None, '長度_m': U['len'][i]},
        'geometry': {'type': 'LineString', 'coordinates': U['geom'][i]}})
(OUT / f'{YID}年度巡檢覆蓋成果_管段明細.geojson').write_text(
    json.dumps({'type': 'FeatureCollection',
                'crs': {'type': 'name', 'properties': {'name': 'urn:ogc:def:crs:OGC:1.3:CRS84'}},
                'features': feats}, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')

cov = sum(U['len'][i] for i in range(N) if ST[i] == 0)
print(f'年度 {YID}｜單元 {N:,}｜總長 {TOT/1000:,.2f} km｜'
      f'已巡檢 {cov/1000:,.2f} km（{cov/TOT*100:.2f}%）')
for f in sorted(OUT.iterdir()):
    print(f'  → out/{f.name}  {f.stat().st_size/1e6:.2f} MB')
